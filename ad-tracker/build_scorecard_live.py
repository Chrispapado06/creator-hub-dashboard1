#!/usr/bin/env python3
"""UNCVRD_Ad_Scorecard_LIVE.xlsx — NO Apps Script version.
  • Feed         : one IMPORTDATA formula pulls daily per-creator data from our server.
  • Manual Spend : boss types OnlyFinder/Guider/Seeker spend here (date, creator, platform, $).
  • Scorecard    : weekly view with Week + Creator pickers, reads Feed + Manual Spend.
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FEED_URL = "https://uncvrd-ad-suite.onrender.com/sheet-feed?key=uncvrd2026"

wb = Workbook()
BAND="5B43F5"; HEAD="1F2430"; WHITE="FFFFFF"; MUTE="6B7280"
AUTO="E8F6EC"; MAN="FCEFD7"; FORM="EEF0F4"; YEL="FFF6CC"
thin=Side(style="thin",color="D5D9E0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
CUR='$#,##0.00'; INT='#,##0'; PCT='0.0%'; MUL='0.00"x"'

def style(c,*,bold=False,size=11,color="111111",fill=None,align="left",fmt=None,wrap=False,bd=True):
    c.font=Font(bold=bold,size=size,color=color)
    c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    if fmt: c.number_format=fmt
    if bd: c.border=border

today=datetime.date(2026,6,12)
sunday=today-datetime.timedelta(days=(today.weekday()+1)%7)

# ---- pull current ad links so we can pre-fill the Link Settings tab ----
def _ad_links():
    import json, urllib.request
    key=json.load(open("/Users/christofispapadopoulos/Downloads/creator-hub-dashboard-main/ad-tracker/data-analyst/credentials.json"))["onlyfans_key"]
    B="https://app.onlyfansapi.com/api"
    def f(p):
        r=urllib.request.Request(B+p,headers={"Authorization":"Bearer "+key,"Accept":"application/json","User-Agent":"UNCVRD/1.0"})
        return json.load(urllib.request.urlopen(r,timeout=45))
    def cls(n):
        n=(n or "").lower()
        if "meta" in n: return "Meta"
        if "guider" in n: return "OnlyGuider"
        if "seeker" in n: return "OnlySeeker"
        if "finder" in n or "search" in n: return "OnlyFinder"
        return ""
    accts=f("/accounts"); accts=accts if isinstance(accts,list) else (accts.get("data") if isinstance(accts.get("data"),list) else (accts.get("data") or {}).get("list") or [])
    out=[]
    for a in accts:
        if not a.get("is_authenticated"): continue
        name=a.get("display_name") or a.get("onlyfans_username") or ""
        off=0
        for _ in range(30):
            d=f("/%s/tracking-links?limit=100&offset=%d"%(a["id"],off)); dd=(d or {}).get("data") or {}; lst=dd.get("list") or []
            for l in lst:
                p=cls(l.get("campaignName"))
                if p: out.append((name, l.get("campaignName") or "", str(l.get("campaignCode") or ""), p))
            if not dd.get("hasMore") or not lst: break
            off+=len(lst)
    return out
try: AD_LINKS=_ad_links()
except Exception: AD_LINKS=[]

# ===================== FEED (auto — one formula, sheet-id aware) =====================
feed=wb.active; feed.title="Feed"
# the &sheet= part tells the server which sheet's Link Settings to read (for overrides)
feed["A1"]=f'=IMPORTDATA("{FEED_URL}&sheet="&\'Link Settings\'!$B$1)'
# normalized date key so the Scorecard can match regardless of how Sheets parses the dates
feed["U1"]="DateKey"
feed["U2"]='=ARRAYFORMULA(IF(LEN(A2:A100000)=0,"",TEXT(A2:A100000,"yyyy-mm-dd")))'
style(feed["U1"],bold=True,fill=HEAD,color=WHITE,align="center")
feed["W1"]="This tab fills itself from the UNCVRD server (history included). Don't type here."
style(feed["W1"],color=MUTE,bd=False)
feed.column_dimensions["A"].width=12; feed.column_dimensions["B"].width=16
feed.column_dimensions["U"].width=12; feed.column_dimensions["W"].width=60

# ===================== MANUAL SPEND =====================
ms=wb.create_sheet("Manual Spend")
ms.merge_cells(start_row=1,start_column=1,end_row=1,end_column=5)
style(ms.cell(1,1,"MANUAL SPEND — type what you spent on OnlyFinder / OnlyGuider / OnlySeeker (no API for those). One row per day per platform."),
      bold=True,color=WHITE,fill=BAND,bd=False)
for j,h in enumerate(["Date","Creator","Platform","Amount ($)"],start=1):
    style(ms.cell(2,j,h),bold=True,color=WHITE,fill=HEAD,align="center")
for r in range(3,103):
    style(ms.cell(r,1),fmt="yyyy-mm-dd",align="center")
    style(ms.cell(r,2),align="center")
    style(ms.cell(r,3),align="center")
    style(ms.cell(r,4),fmt=CUR,align="center",fill=MAN)
dv=DataValidation(type="list",formula1='"OnlyFinder,OnlyGuider,OnlySeeker"',allow_blank=True)
ms.add_data_validation(dv); dv.add("C3:C102")
ms.column_dimensions["A"].width=13; ms.column_dimensions["B"].width=16
ms.column_dimensions["C"].width=14; ms.column_dimensions["D"].width=12
ms.freeze_panes="A3"; ms.sheet_view.showGridLines=False

# ===================== LINK SETTINGS (manual control) =====================
ls=wb.create_sheet("Link Settings")
ls.merge_cells(start_row=1,start_column=3,end_row=1,end_column=5)
style(ls.cell(1,1,"Sheet ID ▸"),bold=True,align="right")
style(ls.cell(1,2,""),fill=YEL,align="center",bold=True)   # boss pastes the sheet's ID here
style(ls.cell(1,3,"⬅ Paste this sheet's ID (the code in its URL between /d/ and /edit) to switch ON manual control. Leave blank = fully automatic."),
      size=10,color=MUTE,bd=False)
ls.merge_cells(start_row=2,start_column=1,end_row=2,end_column=5)
style(ls.cell(2,1,"LINK SETTINGS — leave Override BLANK to auto-detect by name. Or force a platform, or type Ignore to exclude a link."),
      bold=True,color=WHITE,fill=BAND,bd=False)
for j,h in enumerate(["Creator","Tracking link","Code","Auto-detected","Override (your choice)"],start=1):
    style(ls.cell(3,j,h),bold=True,color=WHITE,fill=HEAD,align="center")
r=4
for (creator,link,code,plat) in AD_LINKS:
    style(ls.cell(r,1,creator),align="left")
    style(ls.cell(r,2,link),align="left")
    style(ls.cell(r,3,code),align="center")
    style(ls.cell(r,4,plat),align="center",color=MUTE)
    style(ls.cell(r,5,""),align="center",fill=MAN)
    r+=1
last=max(r,5)
dvl=DataValidation(type="list",formula1='"Meta,OnlyFinder,OnlyGuider,OnlySeeker,Ignore"',allow_blank=True)
ls.add_data_validation(dvl); dvl.add("E4:E%d"%(last+200))
ls.column_dimensions["A"].width=16; ls.column_dimensions["B"].width=30
ls.column_dimensions["C"].width=8; ls.column_dimensions["D"].width=14; ls.column_dimensions["E"].width=20
ls.freeze_panes="A4"; ls.sheet_view.showGridLines=False

# ===================== SCORECARD =====================
sc=wb.create_sheet("Scorecard",0)
FIRST,DAYS=3,7; LAST=FIRST+DAYS-1; AVG,TGT,STAT,SRC,NOTE=10,11,12,13,14
def CL(col,r): return "%s%d"%(get_column_letter(col),r)
CRE="$F$3"

sc.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NOTE)
style(sc.cell(1,1,"UNCVRD — AD SCORECARD (LIVE)"),bold=True,size=18,color=WHITE,fill=BAND,bd=False)
sc.merge_cells(start_row=2,start_column=1,end_row=2,end_column=NOTE)
style(sc.cell(2,1,"Fills itself from the server (Feed tab) — no buttons, no scripts. Pick a Creator and a Week. CPC ≤ $0.70, CVR ≥ 25%."),
      size=10,color=MUTE,bd=False)
style(sc.cell(3,2,"Week of (Sun) ▸"),bold=True,align="right")
style(sc.cell(3,3,sunday),fmt="yyyy-mm-dd",align="center",fill=YEL,bold=True)
style(sc.cell(3,5,"Creator ▸"),bold=True,align="right")
style(sc.cell(3,6,"All"),align="center",fill=YEL,bold=True)
dvc=DataValidation(type="list",formula1='"All,June - Sandra,Antonella,Marissa munoz"',allow_blank=False)
sc.add_data_validation(dvc); dvc.add("F3")
style(sc.cell(3,8,"← pick creator & week"),size=10,color=MUTE,bd=False)

sc.merge_cells(start_row=4,start_column=2,end_row=5,end_column=2)
style(sc.cell(4,2,"METRIC"),bold=True,color=WHITE,fill=HEAD)
for i in range(DAYS):
    col=FIRST+i
    style(sc.cell(5,col,f"=$C$3+{i}"),color=WHITE,fill=HEAD,align="center",fmt="mmm d")
    style(sc.cell(4,col,f'=TEXT({CL(col,5)},"ddd")'),bold=True,color=WHITE,fill=HEAD,align="center")
for col,label in [(AVG,"Avg/day"),(TGT,"Target"),(STAT,"Status"),(SRC,"Source"),(NOTE,"Notes")]:
    sc.merge_cells(start_row=4,start_column=col,end_row=5,end_column=col)
    style(sc.cell(4,col,label),bold=True,color=WHITE,fill=HEAD,align="center")

def band(r,t):
    sc.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NOTE)
    style(sc.cell(r,1,t),bold=True,color=WHITE,fill=BAND)

def feedcol(logcol):   # auto metric out of the Feed tab
    return lambda col,r:('=SUMIFS(Feed!$%s:$%s,Feed!$U:$U,TEXT(%s$5,"yyyy-mm-dd"),'
                         'Feed!$B:$B,IF(%s="All","*",%s))'
                         %(logcol,logcol,get_column_letter(col),CRE,CRE))
def manual(plat):      # manual spend out of the Manual Spend tab
    return lambda col,r:("=SUMIFS('Manual Spend'!$D:$D,'Manual Spend'!$A:$A,%s$5,"
                         "'Manual Spend'!$B:$B,IF(%s=\"All\",\"*\",%s),'Manual Spend'!$C:$C,\"%s\")"
                         %(get_column_letter(col),CRE,CRE,plat))
def sum_rows(a,b): return lambda col,r:"=SUM(%s:%s)"%(CL(col,a),CL(col,b))
def ratio(a,b):    return lambda col,r:'=IFERROR(%s/%s,"")'%(CL(col,a),CL(col,b))
def diff(a,b):     return lambda col,r:'=%s-%s'%(CL(col,a),CL(col,b))

def metric(r,name,*,fmt,src,note="",target=None,status=None,day_formula,bold_name=False,raw=False):
    style(sc.cell(r,2,name),bold=bold_name)
    fill=WHITE if raw else FORM
    for col in range(FIRST,LAST+1):
        style(sc.cell(r,col,day_formula(col,r)),align="center",fmt=fmt,fill=fill)
    style(sc.cell(r,AVG,'=IFERROR(AVERAGE(%s:%s),"")'%(CL(FIRST,r),CL(LAST,r))),align="center",fmt=fmt,fill=FORM)
    style(sc.cell(r,TGT,target if target is not None else ""),align="center",fmt=(fmt if target is not None else None),bold=True)
    style(sc.cell(r,STAT,status or ""),align="center")
    style(sc.cell(r,SRC,src),align="center",color=MUTE,size=10)
    style(sc.cell(r,NOTE,note),color=MUTE,size=10,wrap=True)

band(6,"SPEND")
metric(7,"Ad Spend — Meta",fmt=CUR,src="Meta API",note="auto (via server)",day_formula=feedcol("C"),raw=True)
metric(8,"Ad Spend — OnlyFinder",fmt=CUR,src="Manual Spend tab",note="type in Manual Spend",day_formula=manual("OnlyFinder"),raw=True)
metric(9,"Ad Spend — OnlyGuider",fmt=CUR,src="Manual Spend tab",note="type in Manual Spend",day_formula=manual("OnlyGuider"),raw=True)
metric(10,"Ad Spend — OnlySeeker",fmt=CUR,src="Manual Spend tab",note="type in Manual Spend",day_formula=manual("OnlySeeker"),raw=True)
metric(11,"Total Ad Spend",fmt=CUR,src="Formula",day_formula=sum_rows(7,10),bold_name=True)
band(12,"TRAFFIC")
metric(13,"Clicks — Meta",fmt=INT,src="OnlyFans API",day_formula=feedcol("G"),raw=True)
metric(14,"Clicks — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=feedcol("H"),raw=True)
metric(15,"Clicks — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=feedcol("I"),raw=True)
metric(16,"Clicks — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=feedcol("J"),raw=True)
metric(17,"Total Clicks",fmt=INT,src="Formula",day_formula=sum_rows(13,16),bold_name=True)
metric(18,"CPC — Meta",fmt=CUR,src="Formula",day_formula=ratio(7,13),target=0.70,
       status='=IF(N(J18)=0,"",IF(J18<=K18,"✅","⚠️"))',note="≤ $0.70")
metric(19,"CPC — OnlyFinder",fmt=CUR,src="Formula",day_formula=ratio(8,14),target=0.70,
       status='=IF(N(J19)=0,"",IF(J19<=K19,"✅","⚠️"))',note="≤ $0.70")
metric(20,"Blended CPC",fmt=CUR,src="Formula",day_formula=ratio(11,17),target=0.70,
       status='=IF(N(J20)=0,"",IF(J20<=K20,"✅","⚠️"))')
band(21,"FANS (CONVERSIONS)")
metric(22,"New Fans — Meta",fmt=INT,src="OnlyFans API",note="links named meta",day_formula=feedcol("K"),raw=True)
metric(23,"New Fans — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=feedcol("L"),raw=True)
metric(24,"New Fans — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=feedcol("M"),raw=True)
metric(25,"New Fans — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=feedcol("N"),raw=True)
metric(26,"Total New Fans",fmt=INT,src="Formula",day_formula=sum_rows(22,25),bold_name=True)
metric(27,"Subscription CVR — Meta",fmt=PCT,src="Formula",day_formula=ratio(22,13),target=0.25,
       status='=IF(N(J27)=0,"",IF(J27>=K27,"✅","⚠️"))',note="≥ 25% else split-test")
metric(28,"Subscription CVR — Blended",fmt=PCT,src="Formula",day_formula=ratio(26,17),target=0.25,
       status='=IF(N(J28)=0,"",IF(J28>=K28,"✅","⚠️"))')
metric(29,"Cost per Fan (CAC)",fmt=CUR,src="Formula",day_formula=ratio(11,26))
band(30,"REVENUE")
metric(31,"Revenue",fmt=CUR,src="OnlyFans API",day_formula=feedcol("O"),raw=True)
metric(32,"Revenue per Fan (LTV)",fmt=CUR,src="Formula",day_formula=ratio(31,26))
metric(33,"ROAS",fmt=MUL,src="Formula",day_formula=ratio(31,11),target=2.00,
       status='=IF(N(J33)=0,"",IF(J33>=K33,"✅","⚠️"))',note="≥2x scale · <1x cut")
metric(34,"Profit",fmt=CUR,src="Formula",day_formula=diff(31,11),bold_name=True)

sc.column_dimensions["A"].width=2; sc.column_dimensions["B"].width=26
for col in range(FIRST,LAST+1): sc.column_dimensions[get_column_letter(col)].width=9.5
for col,w in [(AVG,9.5),(TGT,9),(STAT,8),(SRC,15),(NOTE,24)]:
    sc.column_dimensions[get_column_letter(col)].width=w
sc.freeze_panes="C6"; sc.sheet_view.showGridLines=False; sc.row_dimensions[1].height=26

out="/Users/christofispapadopoulos/Downloads/creator-hub-dashboard-main/ad-tracker/UNCVRD_Ad_Scorecard_LIVE.xlsx"
wb.save(out); print("wrote",out)
