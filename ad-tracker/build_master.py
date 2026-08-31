#!/usr/bin/env python3
"""UNCVRD_Ad_Suite.xlsx — the complete sheet, import-ready.

Tabs (left to right):
  Start Here      — one-time setup (paste sheet ID) + guide
  Compare         — pick a date range, compare platforms side by side (boss's simple view)
  Daily Breakdown — flat table: one row per day+platform, all metrics (IMPORTDATA /sheet-flat)
  Scorecard       — weekly per-creator detail (Week + Creator pickers), reads Feed
  Manual Spend    — type OnlyFinder/Guider/Seeker spend
  Manual Clicks   — type clicks to override the auto number
  Link Settings   — force a platform / ignore a link
  Feed            — raw auto data from the server (don't touch)

Everything fills itself from the UNCVRD server. The only typing the boss ever does
is in Manual Spend / Manual Clicks / Link Settings.
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = "https://uncvrd-ad-suite.onrender.com"
KEY = "uncvrd2026"
SID = "'Start Here'!$B$2"     # the cell where the boss pastes this sheet's ID

BAND="5B43F5"; HEAD="1F2430"; WHITE="FFFFFF"; MUTE="6B7280"
AUTO="E8F6EC"; MAN="FCEFD7"; FORM="EEF0F4"; YEL="FFF6CC"; CARD="F6F7FB"
thin=Side(style="thin",color="D5D9E0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
CUR='$#,##0.00'; INT='#,##0'; PCT='0.0%'; MUL='0.00"x"'

def style(c,*,bold=False,size=11,color="111111",fill=None,align="left",fmt=None,wrap=False,bd=True):
    c.font=Font(bold=bold,size=size,color=color)
    c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    if fmt: c.number_format=fmt
    if bd: c.border=border

today=datetime.date(2026,6,15)
sunday=today-datetime.timedelta(days=(today.weekday()+1)%7)

# ---- pull current ad links to pre-fill Link Settings ----
def _ad_links():
    import json, urllib.request
    key=json.load(open("/Users/christofispapadopoulos/Downloads/creator-hub-dashboard-main/ad-tracker/data-analyst/credentials.json"))["onlyfans_key"]
    B="https://app.onlyfansapi.com/api"
    def f(p):
        r=urllib.request.Request(B+p,headers={"Authorization":"Bearer "+key,"Accept":"application/json","User-Agent":"UNCVRD/1.0"})
        return json.load(urllib.request.urlopen(r,timeout=45))
    def cls(n):
        n=(n or "").lower()
        if "meta" in n: return "Meta"
        if "guider" in n: return "OnlyGuider"
        if "seeker" in n: return "OnlySeeker"
        if "finder" in n or "search" in n: return "OnlyFinder"
        return ""
    accts=f("/accounts"); accts=accts if isinstance(accts,list) else (accts.get("data") if isinstance(accts.get("data"),list) else (accts.get("data") or {}).get("list") or [])
    out=[]
    for a in accts:
        if not a.get("is_authenticated"): continue
        name=a.get("display_name") or a.get("onlyfans_username") or ""
        off=0
        for _ in range(30):
            d=f("/%s/tracking-links?limit=100&offset=%d"%(a["id"],off)); dd=(d or {}).get("data") or {}; lst=dd.get("list") or []
            for l in lst:
                p=cls(l.get("campaignName"))
                if p: out.append((name, l.get("campaignName") or "", str(l.get("campaignCode") or ""), p))
            if not dd.get("hasMore") or not lst: break
            off+=len(lst)
    return out
try: AD_LINKS=_ad_links()
except Exception: AD_LINKS=[]
CREATORS=[]
for (cr,_l,_c,_p) in AD_LINKS:
    if cr not in CREATORS: CREATORS.append(cr)

PLATS=["Meta","OnlyFinder","OnlyGuider","OnlySeeker"]
wb=Workbook(); wb.remove(wb.active)

# ============================================================ START HERE
sh=wb.create_sheet("Start Here")
sh.merge_cells("A1:F1")
style(sh.cell(1,1,"UNCVRD — AD SUITE"),bold=True,size=20,color=WHITE,fill=BAND,bd=False)
# ONE input cell only: B2 is the sheet-ID anchor (SID = 'Start Here'!$B$2). Don't move it.
style(sh.cell(2,1,"① Paste this sheet's ID here ▸"),bold=True,align="right")
style(sh.cell(2,2,""),fill=YEL,bold=True,align="center")
sh.merge_cells("C2:F2")
style(sh.cell(2,3,"Copy the long code in this sheet's URL (between /d/ and /edit) and paste it in the yellow box. This switches ON the Manual Spend / Clicks / Link Settings tabs — the auto tabs already fill themselves."),
      size=10,color=MUTE,wrap=True,bd=False)
# guide
style(sh.cell(4,1,"WHAT EACH TAB IS FOR"),bold=True,size=12,fill=BAND,color=WHITE,bd=False)
sh.merge_cells("A4:F4")
guide=[("Compare","Pick a date range → see every platform side by side. Your quick comparison view.","read"),
       ("Daily Breakdown","One row per day + platform, every metric. Newest on top. Auto.","read"),
       ("Scorecard","Weekly detail — pick a creator + week. Auto.","read"),
       ("Manual Spend","TYPE here: what you spent on OnlyFinder / OnlyGuider / OnlySeeker.","type"),
       ("Manual Clicks","TYPE here only to override a day's clicks (e.g. OnlyFinder's billed number).","type"),
       ("Link Settings","Force a link's platform, or 'Ignore' it. Blank = automatic.","type"),
       ("Feed","Raw auto data from the server. DON'T TOUCH.","auto")]
style(sh.cell(5,1,"Tab"),bold=True,fill=HEAD,color=WHITE,align="center")
style(sh.cell(5,2,"What it's for"),bold=True,fill=HEAD,color=WHITE)
sh.merge_cells("B5:E5")
style(sh.cell(5,6,"You…"),bold=True,fill=HEAD,color=WHITE,align="center")
r=6
for (tab,desc,kind) in guide:
    style(sh.cell(r,1,tab),bold=True)
    sh.merge_cells(start_row=r,start_column=2,end_row=r,end_column=5)
    style(sh.cell(r,2,desc),wrap=True)
    tag={"read":"just read","type":"TYPE","auto":"don't touch"}[kind]
    fl={"read":CARD,"type":MAN,"auto":FORM}[kind]
    style(sh.cell(r,6,tag),align="center",fill=fl)
    r+=1
sh.column_dimensions["A"].width=18
for col in ["B","C","D","E"]: sh.column_dimensions[col].width=20
sh.column_dimensions["F"].width=12
sh.sheet_view.showGridLines=False; sh.row_dimensions[1].height=30

# ============================================================ FEED (auto)
feed=wb.create_sheet("Feed")
feed["A1"]=f'=IFERROR(IMPORTDATA("{BASE}/sheet-feed?key={KEY}&sheet="&{SID}),"")'
feed["U1"]="DateKey"
feed["U2"]='=ARRAYFORMULA(IF(LEN(A2:A100000)=0,"",TEXT(A2:A100000,"yyyy-mm-dd")))'
style(feed["U1"],bold=True,fill=HEAD,color=WHITE,align="center")
feed["W1"]="Raw auto data from the UNCVRD server. Don't type here."
style(feed["W1"],color=MUTE,bd=False)
feed.column_dimensions["A"].width=12; feed.column_dimensions["B"].width=16
feed.column_dimensions["U"].width=12; feed.column_dimensions["W"].width=50

# ============================================================ DAILY BREAKDOWN (auto, flat)
db=wb.create_sheet("Daily Breakdown")
db["A1"]=f'=IFERROR(IMPORTDATA("{BASE}/sheet-flat?key={KEY}&sheet="&{SID}),"")'
# text date-key (col N) so the Compare tab matches text-to-text (robust, locale-proof)
db["N1"]="DateKey"
db["N2"]='=ARRAYFORMULA(IF(LEN(A2:A100000)=0,"",TEXT(A2:A100000,"yyyy-mm-dd")))'
style(db["N1"],bold=True,fill=HEAD,color=WHITE,align="center")
db["P1"]="One row per day + platform · newest on top · fills itself. (Spend columns need Manual Spend / Meta token.)"
style(db["P1"],color=MUTE,bd=False)
for col,w in [("A",12),("B",12),("C",9),("D",8),("E",10),("F",12),("G",9),("H",9),("I",14),("J",13),("K",9),("L",13)]:
    db.column_dimensions[col].width=w
db.freeze_panes="A2"; db.sheet_view.showGridLines=False

# ============================================================ COMPARE (simple, date-range)
cmp=wb.create_sheet("Compare")
def C(col,r): return "%s%d"%(get_column_letter(col),r)
cmp.merge_cells("A1:L1"); style(cmp.cell(1,1,"COMPARE PLATFORMS"),bold=True,size=18,color=WHITE,fill=BAND,bd=False)
style(cmp.cell(2,1,"Pick a date range — every platform's totals line up side by side. Reads from Daily Breakdown."),size=10,color=MUTE,bd=False)
cmp.merge_cells("A2:L2")
style(cmp.cell(3,1,"From ▸"),bold=True,align="right"); style(cmp.cell(3,2,today-datetime.timedelta(days=6)),fmt="yyyy-mm-dd",fill=YEL,bold=True,align="center")
style(cmp.cell(3,3,"To ▸"),bold=True,align="right"); style(cmp.cell(3,4,today),fmt="yyyy-mm-dd",fill=YEL,bold=True,align="center")
style(cmp.cell(3,6,"← change the range to compare any period"),size=10,color=MUTE,bd=False)
hdrs=["Platform","Clicks","Fans","Spend","CPC","Cost / Fan","CVR","Revenue","ROAS","Profit"]
for j,h in enumerate(hdrs,start=1): style(cmp.cell(4,j,h),bold=True,fill=HEAD,color=WHITE,align="center")
FROM='">="&TEXT($B$3,"yyyy-mm-dd")'; TO='"<="&TEXT($D$3,"yyyy-mm-dd")'
def db_sumifs(valcol, plat_expr):
    # match the TEXT DateKey (col N) text-to-text over the range; sum the value column
    base=("SUMIFS('Daily Breakdown'!$%s:$%s,'Daily Breakdown'!$N:$N,%s,"
          "'Daily Breakdown'!$N:$N,%s"%(valcol,valcol,FROM,TO))
    if plat_expr is not None:
        base+=",'Daily Breakdown'!$B:$B,%s"%plat_expr
    return base+")"
rows_order=PLATS+["TOTAL"]
r=5
for p in rows_order:
    pe=None if p=="TOTAL" else '"%s"'%p
    style(cmp.cell(r,1,p),bold=(p=="TOTAL"))
    clicks="="+db_sumifs("C",pe); fans="="+db_sumifs("D",pe); spend="="+db_sumifs("E",pe); rev="="+db_sumifs("I",pe)
    style(cmp.cell(r,2,clicks),align="center",fmt=INT)
    style(cmp.cell(r,3,fans),align="center",fmt=INT)
    style(cmp.cell(r,4,spend),align="center",fmt=CUR)
    style(cmp.cell(r,5,'=IFERROR(%s/%s,"")'%(C(4,r),C(2,r))),align="center",fmt=CUR)   # CPC
    style(cmp.cell(r,6,'=IFERROR(%s/%s,"")'%(C(4,r),C(3,r))),align="center",fmt=CUR)   # cost/fan
    style(cmp.cell(r,7,'=IFERROR(%s/%s,"")'%(C(3,r),C(2,r))),align="center",fmt=PCT)   # CVR
    style(cmp.cell(r,8,rev),align="center",fmt=CUR)
    style(cmp.cell(r,9,'=IFERROR(%s/%s,"")'%(C(8,r),C(4,r))),align="center",fmt=MUL)   # ROAS
    style(cmp.cell(r,10,'=%s-%s'%(C(8,r),C(4,r))),align="center",fmt=CUR)           # profit
    r+=1
cmp.column_dimensions["A"].width=14
for col in ["B","C","D","E","F","G","H","I","J"]: cmp.column_dimensions[col].width=11
cmp.freeze_panes="A5"; cmp.sheet_view.showGridLines=False; cmp.row_dimensions[1].height=26

# ============================================================ SCORECARD (weekly per-creator)
sc=wb.create_sheet("Scorecard")
FIRST,DAYS=3,7; LAST=FIRST+DAYS-1; AVG,TGT,STAT,SRC,NOTE=10,11,12,13,14
def CL(col,r): return "%s%d"%(get_column_letter(col),r)
CRE="$F$3"
sc.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NOTE)
style(sc.cell(1,1,"UNCVRD — WEEKLY SCORECARD"),bold=True,size=18,color=WHITE,fill=BAND,bd=False)
sc.merge_cells(start_row=2,start_column=1,end_row=2,end_column=NOTE)
style(sc.cell(2,1,"Pick a Creator and a Week. Auto from the server. CPC ≤ $0.70, CVR ≥ 25%."),size=10,color=MUTE,bd=False)
style(sc.cell(3,2,"Week of (Sun) ▸"),bold=True,align="right")
style(sc.cell(3,3,sunday),fmt="yyyy-mm-dd",align="center",fill=YEL,bold=True)
style(sc.cell(3,5,"Creator ▸"),bold=True,align="right")
style(sc.cell(3,6,"All"),align="center",fill=YEL,bold=True)
dvc=DataValidation(type="list",formula1='"%s"'%(",".join(["All"]+CREATORS)),allow_blank=False)
sc.add_data_validation(dvc); dvc.add("F3")
sc.merge_cells(start_row=4,start_column=2,end_row=5,end_column=2)
style(sc.cell(4,2,"METRIC"),bold=True,color=WHITE,fill=HEAD)
for i in range(DAYS):
    col=FIRST+i
    style(sc.cell(5,col,f"=$C$3+{i}"),color=WHITE,fill=HEAD,align="center",fmt="mmm d")
    style(sc.cell(4,col,f'=TEXT({CL(col,5)},"ddd")'),bold=True,color=WHITE,fill=HEAD,align="center")
for col,label in [(AVG,"Avg/day"),(TGT,"Target"),(STAT,"Status"),(SRC,"Source"),(NOTE,"Notes")]:
    sc.merge_cells(start_row=4,start_column=col,end_row=5,end_column=col)
    style(sc.cell(4,col,label),bold=True,color=WHITE,fill=HEAD,align="center")
def band(r,t):
    sc.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NOTE)
    style(sc.cell(r,1,t),bold=True,color=WHITE,fill=BAND)
def feedcol(letter):
    return lambda col,r:('=SUMIFS(Feed!$%s:$%s,Feed!$U:$U,TEXT(%s$5,"yyyy-mm-dd"),'
                         'Feed!$B:$B,IF(%s="All","*",%s))'%(letter,letter,get_column_letter(col),CRE,CRE))
def sum_rows(a,b): return lambda col,r:"=SUM(%s:%s)"%(CL(col,a),CL(col,b))
def ratio(a,b):    return lambda col,r:'=IFERROR(%s/%s,"")'%(CL(col,a),CL(col,b))
def diff(a,b):     return lambda col,r:'=%s-%s'%(CL(col,a),CL(col,b))
def metric(r,name,*,fmt,src,note="",target=None,status=None,day_formula,bold_name=False,raw=False):
    style(sc.cell(r,2,name),bold=bold_name)
    fill=WHITE if raw else FORM
    for col in range(FIRST,LAST+1):
        style(sc.cell(r,col,day_formula(col,r)),align="center",fmt=fmt,fill=fill)
    style(sc.cell(r,AVG,'=IFERROR(AVERAGE(%s:%s),"")'%(CL(FIRST,r),CL(LAST,r))),align="center",fmt=fmt,fill=FORM)
    style(sc.cell(r,TGT,target if target is not None else ""),align="center",fmt=(fmt if target is not None else None),bold=True)
    style(sc.cell(r,STAT,status or ""),align="center")
    style(sc.cell(r,SRC,src),align="center",color=MUTE,size=10)
    style(sc.cell(r,NOTE,note),color=MUTE,size=10,wrap=True)
# Feed letters: C AdSpendMeta, D OF, E OG, F OS | G ClicksMeta H OF I OG J OS | K FansMeta L OF M OG N OS | O Revenue
band(6,"SPEND")
metric(7,"Ad Spend — Meta",fmt=CUR,src="Meta API",note="auto",day_formula=feedcol("C"),raw=True)
metric(8,"Ad Spend — OnlyFinder",fmt=CUR,src="Manual Spend tab",day_formula=feedcol("D"),raw=True)
metric(9,"Ad Spend — OnlyGuider",fmt=CUR,src="Manual Spend tab",day_formula=feedcol("E"),raw=True)
metric(10,"Ad Spend — OnlySeeker",fmt=CUR,src="Manual Spend tab",day_formula=feedcol("F"),raw=True)
metric(11,"Total Ad Spend",fmt=CUR,src="Formula",day_formula=sum_rows(7,10),bold_name=True)
band(12,"TRAFFIC")
metric(13,"Clicks — Meta",fmt=INT,src="OnlyFans API",day_formula=feedcol("G"),raw=True)
metric(14,"Clicks — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=feedcol("H"),raw=True)
metric(15,"Clicks — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=feedcol("I"),raw=True)
metric(16,"Clicks — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=feedcol("J"),raw=True)
metric(17,"Total Clicks",fmt=INT,src="Formula",day_formula=sum_rows(13,16),bold_name=True)
metric(18,"CPC — Meta",fmt=CUR,src="Formula",day_formula=ratio(7,13),target=0.70,
       status='=IF(N(J18)=0,"",IF(J18<=K18,"✅","⚠️"))',note="≤ $0.70")
metric(19,"CPC — OnlyFinder",fmt=CUR,src="Formula",day_formula=ratio(8,14),target=0.70,
       status='=IF(N(J19)=0,"",IF(J19<=K19,"✅","⚠️"))',note="≤ $0.70")
metric(20,"Blended CPC",fmt=CUR,src="Formula",day_formula=ratio(11,17),target=0.70,
       status='=IF(N(J20)=0,"",IF(J20<=K20,"✅","⚠️"))')
band(21,"FANS (CONVERSIONS)")
metric(22,"New Fans — Meta",fmt=INT,src="OnlyFans API",note="links named meta",day_formula=feedcol("K"),raw=True)
metric(23,"New Fans — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=feedcol("L"),raw=True)
metric(24,"New Fans — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=feedcol("M"),raw=True)
metric(25,"New Fans — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=feedcol("N"),raw=True)
metric(26,"Total New Fans",fmt=INT,src="Formula",day_formula=sum_rows(22,25),bold_name=True)
metric(27,"Subscription CVR — Meta",fmt=PCT,src="Formula",day_formula=ratio(22,13),target=0.25,
       status='=IF(N(J27)=0,"",IF(J27>=K27,"✅","⚠️"))',note="≥ 25% else split-test")
metric(28,"Subscription CVR — Blended",fmt=PCT,src="Formula",day_formula=ratio(26,17),target=0.25,
       status='=IF(N(J28)=0,"",IF(J28>=K28,"✅","⚠️"))')
metric(29,"Cost per Fan (CAC)",fmt=CUR,src="Formula",day_formula=ratio(11,26))
band(30,"REVENUE")
metric(31,"Revenue",fmt=CUR,src="OnlyFans API",day_formula=feedcol("O"),raw=True)
metric(32,"Revenue per Fan (LTV)",fmt=CUR,src="Formula",day_formula=ratio(31,26))
metric(33,"ROAS",fmt=MUL,src="Formula",day_formula=ratio(31,11),target=2.00,
       status='=IF(N(J33)=0,"",IF(J33>=K33,"✅","⚠️"))',note="≥2x scale · <1x cut")
metric(34,"Profit",fmt=CUR,src="Formula",day_formula=diff(31,11),bold_name=True)
sc.column_dimensions["A"].width=2; sc.column_dimensions["B"].width=26
for col in range(FIRST,LAST+1): sc.column_dimensions[get_column_letter(col)].width=9.5
for col,w in [(AVG,9.5),(TGT,9),(STAT,8),(SRC,15),(NOTE,22)]:
    sc.column_dimensions[get_column_letter(col)].width=w
sc.freeze_panes="C6"; sc.sheet_view.showGridLines=False; sc.row_dimensions[1].height=26

# ============================================================ MANUAL SPEND
ms=wb.create_sheet("Manual Spend")
ms.merge_cells("A1:D1")
style(ms.cell(1,1,"MANUAL SPEND — type what you spent on OnlyFinder / OnlyGuider / OnlySeeker (no API). One row per day."),bold=True,color=WHITE,fill=BAND,bd=False)
for j,h in enumerate(["Date","Creator","Platform","Amount ($)"],start=1):
    style(ms.cell(2,j,h),bold=True,color=WHITE,fill=HEAD,align="center")
for r in range(3,203):
    style(ms.cell(r,1),fmt="yyyy-mm-dd",align="center")
    style(ms.cell(r,2),align="center"); style(ms.cell(r,3),align="center")
    style(ms.cell(r,4),fmt=CUR,align="center",fill=MAN)
dv=DataValidation(type="list",formula1='"OnlyFinder,OnlyGuider,OnlySeeker"',allow_blank=True); ms.add_data_validation(dv); dv.add("C3:C202")
if CREATORS:
    dvcr=DataValidation(type="list",formula1='"%s"'%(",".join(CREATORS)),allow_blank=True); ms.add_data_validation(dvcr); dvcr.add("B3:B202")
ms.column_dimensions["A"].width=13; ms.column_dimensions["B"].width=16; ms.column_dimensions["C"].width=14; ms.column_dimensions["D"].width=12
ms.freeze_panes="A3"; ms.sheet_view.showGridLines=False

# ============================================================ MANUAL CLICKS
mc=wb.create_sheet("Manual Clicks")
mc.merge_cells("A1:D1")
style(mc.cell(1,1,"MANUAL CLICKS — type a number ONLY to override the auto clicks for a day/creator/platform."),bold=True,color=WHITE,fill=BAND,bd=False)
for j,h in enumerate(["Date","Creator","Platform","Clicks"],start=1):
    style(mc.cell(2,j,h),bold=True,color=WHITE,fill=HEAD,align="center")
for r in range(3,203):
    style(mc.cell(r,1),fmt="yyyy-mm-dd",align="center"); style(mc.cell(r,2),align="center"); style(mc.cell(r,3),align="center")
    style(mc.cell(r,4),fmt=INT,align="center",fill=MAN)
dv2=DataValidation(type="list",formula1='"Meta,OnlyFinder,OnlyGuider,OnlySeeker"',allow_blank=True); mc.add_data_validation(dv2); dv2.add("C3:C202")
if CREATORS:
    dvcr2=DataValidation(type="list",formula1='"%s"'%(",".join(CREATORS)),allow_blank=True); mc.add_data_validation(dvcr2); dvcr2.add("B3:B202")
mc.column_dimensions["A"].width=13; mc.column_dimensions["B"].width=16; mc.column_dimensions["C"].width=14; mc.column_dimensions["D"].width=10
mc.freeze_panes="A3"; mc.sheet_view.showGridLines=False

# ============================================================ LINK SETTINGS
ls=wb.create_sheet("Link Settings")
ls.merge_cells("A1:E1")
style(ls.cell(1,1,"LINK SETTINGS — blank Override = auto by name. Force a platform, or type Ignore to drop a link."),bold=True,color=WHITE,fill=BAND,bd=False)
for j,h in enumerate(["Creator","Tracking link","Code","Auto-detected","Override"],start=1):
    style(ls.cell(2,j,h),bold=True,color=WHITE,fill=HEAD,align="center")
r=3
for (cr,link,code,plat) in AD_LINKS:
    style(ls.cell(r,1,cr)); style(ls.cell(r,2,link)); style(ls.cell(r,3,code),align="center")
    style(ls.cell(r,4,plat),align="center",color=MUTE); style(ls.cell(r,5,""),align="center",fill=MAN)
    r+=1
dvl=DataValidation(type="list",formula1='"Meta,OnlyFinder,OnlyGuider,OnlySeeker,Ignore"',allow_blank=True); ls.add_data_validation(dvl); dvl.add("E3:E%d"%(max(r,3)+200))
ls.column_dimensions["A"].width=16; ls.column_dimensions["B"].width=30; ls.column_dimensions["C"].width=8; ls.column_dimensions["D"].width=14; ls.column_dimensions["E"].width=18
ls.freeze_panes="A3"; ls.sheet_view.showGridLines=False

# ---- order tabs: Start Here, Compare, Daily Breakdown, Scorecard, Manual Spend, Manual Clicks, Link Settings, Feed ----
order=["Start Here","Compare","Daily Breakdown","Scorecard","Manual Spend","Manual Clicks","Link Settings","Feed"]
wb._sheets.sort(key=lambda s: order.index(s.title))

out="/Users/christofispapadopoulos/Downloads/creator-hub-dashboard-main/ad-tracker/UNCVRD_Ad_Suite.xlsx"
wb.save(out); print("wrote",out,"| tabs:",[s.title for s in wb._sheets])
