#!/usr/bin/env python3
"""UNCVRD_Ad_Scorecard.xlsx — per-creator version.
  • Daily Log : one row per (day, creator), forever. Has a Creator column.
  • Scorecard : weekly view with a Week picker AND a Creator picker (or "All").
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
BAND="5B43F5"; HEAD="1F2430"; WHITE="FFFFFF"; MUTE="6B7280"
AUTO="E8F6EC"; MAN="FCEFD7"; FORM="EEF0F4"; YEL="FFF6CC"
thin=Side(style="thin",color="D5D9E0"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
CUR='$#,##0.00'; INT='#,##0'; PCT='0.0%'; MUL='0.00"x"'

def style(c,*,bold=False,size=11,color="111111",fill=None,align="left",fmt=None,wrap=False,bd=True):
    c.font=Font(bold=bold,size=size,color=color)
    c.alignment=Alignment(horizontal=align,vertical="center",wrap_text=wrap)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    if fmt: c.number_format=fmt
    if bd: c.border=border

today=datetime.date(2026,6,11)
sunday=today-datetime.timedelta(days=(today.weekday()+1)%7)

# ===================== DAILY LOG (Creator column = B) =====================
log=wb.active; log.title="Daily Log"
LOG_COLS=[("Date","date",None),("Creator","key",None),
    ("Ad Spend Meta","auto",CUR),("Ad Spend OnlyFinder","man",CUR),
    ("Ad Spend OnlyGuider","man",CUR),("Ad Spend OnlySeeker","man",CUR),
    ("Clicks Meta","auto",INT),("Clicks OnlyFinder","auto",INT),
    ("Clicks OnlyGuider","auto",INT),("Clicks OnlySeeker","auto",INT),
    ("Fans Meta","auto",INT),("Fans OnlyFinder","auto",INT),
    ("Fans OnlyGuider","auto",INT),("Fans OnlySeeker","auto",INT),
    ("Revenue","auto",CUR),
    ("Total Spend","form",CUR),("Total Fans","form",INT),("ROAS","form",MUL),("Profit","form",CUR)]
log.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(LOG_COLS))
style(log.cell(1,1,"DAILY LOG  —  one row per day PER CREATOR. Amber = type it in; green = the script fills it."),
      bold=True,color=WHITE,fill=BAND,bd=False)
for j,(name,kind,fmt) in enumerate(LOG_COLS,start=1):
    fill={"date":HEAD,"key":HEAD,"auto":AUTO,"man":MAN,"form":FORM}[kind]
    fg=WHITE if kind in("date","key") else "111111"
    style(log.cell(2,j,name),bold=True,color=fg,fill=fill,align="center",wrap=True)
log.column_dimensions["A"].width=12; log.column_dimensions["B"].width=16
for j in range(3,len(LOG_COLS)+1): log.column_dimensions[get_column_letter(j)].width=11
log.freeze_panes="C3"; log.sheet_view.showGridLines=False; log.row_dimensions[2].height=30

# ===================== SCORECARD =====================
sc=wb.create_sheet("Scorecard",0)
FIRST,DAYS=3,7; LAST=FIRST+DAYS-1; AVG,TGT,STAT,SRC,NOTE=10,11,12,13,14
def CL(col,r): return "%s%d"%(get_column_letter(col),r)
CREATOR_CELL="$F$3"   # the creator picker

sc.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NOTE)
style(sc.cell(1,1,"UNCVRD — AD SCORECARD"),bold=True,size=18,color=WHITE,fill=BAND,bd=False)
sc.merge_cells(start_row=2,start_column=1,end_row=2,end_column=NOTE)
style(sc.cell(2,1,"Pick a Creator and a Week below. Pulls from the Daily Log. CPC ≤ $0.70, CVR ≥ 25%."),
      size=10,color=MUTE,bd=False)
# selectors on row 3
style(sc.cell(3,2,"Week of (Sun) ▸"),bold=True,align="right")
style(sc.cell(3,3,sunday),fmt="yyyy-mm-dd",align="center",fill=YEL,bold=True)
style(sc.cell(3,5,"Creator ▸"),bold=True,align="right")
style(sc.cell(3,6,"All"),align="center",fill=YEL,bold=True)
style(sc.cell(3,8,"← pick creator & week"),size=10,color=MUTE,bd=False)

sc.merge_cells(start_row=4,start_column=2,end_row=5,end_column=2)
style(sc.cell(4,2,"METRIC"),bold=True,color=WHITE,fill=HEAD)
for i in range(DAYS):
    col=FIRST+i
    style(sc.cell(5,col,f"=$C$3+{i}"),color=WHITE,fill=HEAD,align="center",fmt="mmm d")
    style(sc.cell(4,col,f'=TEXT({CL(col,5)},"ddd")'),bold=True,color=WHITE,fill=HEAD,align="center")
for col,label in [(AVG,"Avg/day"),(TGT,"Target"),(STAT,"Status"),(SRC,"Source"),(NOTE,"Notes")]:
    sc.merge_cells(start_row=4,start_column=col,end_row=5,end_column=col)
    style(sc.cell(4,col,label),bold=True,color=WHITE,fill=HEAD,align="center")

def band(r,t):
    sc.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NOTE)
    style(sc.cell(r,1,t),bold=True,color=WHITE,fill=BAND)

def sumifs(logcol):
    return lambda col,r:("=SUMIFS('Daily Log'!$%s:$%s,'Daily Log'!$A:$A,%s$5,"
                         "'Daily Log'!$B:$B,IF(%s=\"All\",\"*\",%s))"
                         %(logcol,logcol,get_column_letter(col),CREATOR_CELL,CREATOR_CELL))
def sum_rows(a,b): return lambda col,r:"=SUM(%s:%s)"%(CL(col,a),CL(col,b))
def ratio(a,b):    return lambda col,r:'=IFERROR(%s/%s,"")'%(CL(col,a),CL(col,b))
def diff(a,b):     return lambda col,r:'=%s-%s'%(CL(col,a),CL(col,b))

def metric(r,name,*,fmt,src,note="",target=None,status=None,day_formula,bold_name=False,raw=False):
    style(sc.cell(r,2,name),bold=bold_name)
    fill=WHITE if raw else FORM
    for col in range(FIRST,LAST+1):
        style(sc.cell(r,col,day_formula(col,r)),align="center",fmt=fmt,fill=fill)
    style(sc.cell(r,AVG,'=IFERROR(AVERAGE(%s:%s),"")'%(CL(FIRST,r),CL(LAST,r))),align="center",fmt=fmt,fill=FORM)
    style(sc.cell(r,TGT,target if target is not None else ""),align="center",fmt=(fmt if target is not None else None),bold=True)
    style(sc.cell(r,STAT,status or ""),align="center")
    style(sc.cell(r,SRC,src),align="center",color=MUTE,size=10)
    style(sc.cell(r,NOTE,note),color=MUTE,size=10,wrap=True)

band(6,"SPEND")
metric(7,"Ad Spend — Meta",fmt=CUR,src="Meta API",note="auto",day_formula=sumifs("C"),raw=True)
metric(8,"Ad Spend — OnlyFinder",fmt=CUR,src="Manual",note="type in Daily Log",day_formula=sumifs("D"),raw=True)
metric(9,"Ad Spend — OnlyGuider",fmt=CUR,src="Manual",note="type in Daily Log",day_formula=sumifs("E"),raw=True)
metric(10,"Ad Spend — OnlySeeker",fmt=CUR,src="Manual",note="type in Daily Log",day_formula=sumifs("F"),raw=True)
metric(11,"Total Ad Spend",fmt=CUR,src="Formula",day_formula=sum_rows(7,10),bold_name=True)
band(12,"TRAFFIC")
metric(13,"Clicks — Meta",fmt=INT,src="OnlyFans API",day_formula=sumifs("G"),raw=True)
metric(14,"Clicks — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=sumifs("H"),raw=True)
metric(15,"Clicks — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=sumifs("I"),raw=True)
metric(16,"Clicks — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=sumifs("J"),raw=True)
metric(17,"Total Clicks",fmt=INT,src="Formula",day_formula=sum_rows(13,16),bold_name=True)
metric(18,"CPC — Meta",fmt=CUR,src="Formula",day_formula=ratio(7,13),target=0.70,
       status='=IF(N(J18)=0,"",IF(J18<=K18,"✅","⚠️"))',note="≤ $0.70")
metric(19,"CPC — OnlyFinder",fmt=CUR,src="Formula",day_formula=ratio(8,14),target=0.70,
       status='=IF(N(J19)=0,"",IF(J19<=K19,"✅","⚠️"))',note="≤ $0.70")
metric(20,"Blended CPC",fmt=CUR,src="Formula",day_formula=ratio(11,17),target=0.70,
       status='=IF(N(J20)=0,"",IF(J20<=K20,"✅","⚠️"))')
band(21,"FANS (CONVERSIONS)")
metric(22,"New Fans — Meta",fmt=INT,src="OnlyFans API",note="links tagged Meta",day_formula=sumifs("K"),raw=True)
metric(23,"New Fans — OnlyFinder",fmt=INT,src="OnlyFans API",day_formula=sumifs("L"),raw=True)
metric(24,"New Fans — OnlyGuider",fmt=INT,src="OnlyFans API",day_formula=sumifs("M"),raw=True)
metric(25,"New Fans — OnlySeeker",fmt=INT,src="OnlyFans API",day_formula=sumifs("N"),raw=True)
metric(26,"Total New Fans",fmt=INT,src="Formula",day_formula=sum_rows(22,25),bold_name=True)
metric(27,"Subscription CVR — Meta",fmt=PCT,src="Formula",day_formula=ratio(22,13),target=0.25,
       status='=IF(N(J27)=0,"",IF(J27>=K27,"✅","⚠️"))',note="≥ 25% else split-test")
metric(28,"Subscription CVR — Blended",fmt=PCT,src="Formula",day_formula=ratio(26,17),target=0.25,
       status='=IF(N(J28)=0,"",IF(J28>=K28,"✅","⚠️"))')
metric(29,"Cost per Fan (CAC)",fmt=CUR,src="Formula",day_formula=ratio(11,26))
band(30,"REVENUE")
metric(31,"Revenue",fmt=CUR,src="OnlyFans API",day_formula=sumifs("O"),raw=True)
metric(32,"Revenue per Fan (LTV)",fmt=CUR,src="Formula",day_formula=ratio(31,26))
metric(33,"ROAS",fmt=MUL,src="Formula",day_formula=ratio(31,11),target=2.00,
       status='=IF(N(J33)=0,"",IF(J33>=K33,"✅","⚠️"))',note="≥2x scale · <1x cut")
metric(34,"Profit",fmt=CUR,src="Formula",day_formula=diff(31,11),bold_name=True)

sc.column_dimensions["A"].width=2; sc.column_dimensions["B"].width=26
for col in range(FIRST,LAST+1): sc.column_dimensions[get_column_letter(col)].width=9.5
for col,w in [(AVG,9.5),(TGT,9),(STAT,8),(SRC,13),(NOTE,22)]:
    sc.column_dimensions[get_column_letter(col)].width=w
sc.freeze_panes="C6"; sc.sheet_view.showGridLines=False; sc.row_dimensions[1].height=26

out="/Users/christofispapadopoulos/Downloads/creator-hub-dashboard-main/ad-tracker/UNCVRD_Ad_Scorecard.xlsx"
wb.save(out); print("wrote",out)
